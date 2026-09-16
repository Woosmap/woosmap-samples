"""Read rows from a CSV file, an Excel workbook or a published Google Sheet."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import requests
from openpyxl import load_workbook

Row = dict[str, str]


def read_source(source: str, sheet: str | None, session: requests.Session) -> list[Row]:
    if urlparse(source).scheme in ("http", "https"):
        return read_google_sheet(source, session)
    path = Path(source)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return read_xlsx_file(path, sheet)
    return read_csv_file(path)


def read_csv_text(text: str) -> list[Row]:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return [{k.strip(): (v or "").strip() for k, v in row.items() if k} for row in reader]


def read_csv_file(path: Path) -> list[Row]:
    return read_csv_text(path.read_text(encoding="utf-8-sig"))


def read_xlsx_file(path: Path, sheet: str | None) -> list[Row]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    return [
        {key: cell_text(value) for key, value in zip(header, row, strict=False) if key}
        for row in rows
        if any(value is not None for value in row)
    ]


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def google_sheet_export_url(url: str) -> str:
    match = re.search(r"/spreadsheets/d/([\w-]+)", url)
    if not match:
        raise ValueError(f"Not a Google Sheets URL: {url}")
    gid_match = re.search(r"[#&?]gid=(\d+)", url)
    gid = f"&gid={gid_match.group(1)}" if gid_match else ""
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=csv{gid}"


def read_google_sheet(url: str, session: requests.Session) -> list[Row]:
    response = session.get(google_sheet_export_url(url), timeout=30)
    response.raise_for_status()
    return read_csv_text(response.content.decode("utf-8-sig"))
